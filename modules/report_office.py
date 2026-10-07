"""
The report as a Word document (.docx) and a PowerPoint deck (.pptx) -- Streamlit-free.

Both render from the same `Section`s as every other format, with equations typeset to images by the
exporter's matplotlib mathtext path (Word and PowerPoint have no way to take a LaTeX string directly
short of Office Math markup, and converting LaTeX to that reliably is its own project). An equation that
mathtext cannot typeset falls back to its LaTeX as monospaced text rather than being dropped.

Word: headings, paragraphs, real tables, bullet lists, equations as inline pictures, plots as pictures.

PowerPoint, built for teaching and review rather than as a wall of text:
  * a title slide, then one slide per SECTION, spilling onto a "(cont.)" slide when it would overflow;
  * the step-by-step section is instead ONE SLIDE PER STEP -- the step's description as the heading, the
    equation large, the explanation beneath, and the explanation again in the speaker notes -- so a
    worked solution can be walked through one step at a time (it is the layout for tutor and worksheet use);
  * a plot you marked is a full-slide picture with its caption.
Live-only figures (the HTML report's interactive charts) have no image and are left out of both.
"""
import io
import math
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.util import Inches as PInches, Pt as PPt

from modules.exporter import _render_latex_image
from modules.report_content import REPORT_TITLE, Section, generated_stamp

_EQ_HEIGHT_IN = {"plain": 0.38, "step": 0.32, "matrix": 0.8, "raw": 0.3}
_MAX_EQ_WIDTH_IN = 6.3


def _eq_picture(latex: str, style: str) -> tuple[io.BytesIO, float, float] | None:
    """(PNG bytes, width in, height in) of a typeset equation sized by its style, or None if it can't be typeset."""
    if style == "raw":
        return None
    img = _render_latex_image(latex, fontsize=13)
    if img is None:
        return None
    data = img.getvalue()
    try:
        w_px, h_px = Image.open(io.BytesIO(data)).size
    except Exception:  # noqa: BLE001
        return None
    height = _EQ_HEIGHT_IN.get(style, 0.38)
    width = height * w_px / h_px
    if width > _MAX_EQ_WIDTH_IN:
        width, height = _MAX_EQ_WIDTH_IN, _MAX_EQ_WIDTH_IN * h_px / w_px
    return io.BytesIO(data), width, height


# --------------------------------------------------------------------------- Word

def render_docx(sections: list[Section], title: str = REPORT_TITLE) -> bytes:
    doc = Document()
    doc.add_heading(title, 0)
    stamp = doc.add_paragraph()
    run = stamp.add_run(generated_stamp())
    run.font.size, run.font.color.rgb = Pt(9), RGBColor(0x77, 0x77, 0x77)

    for sec in sections:
        doc.add_heading(sec.title, 1)
        for b in sec.blocks:
            if b.kind == "text":
                run = doc.add_paragraph().add_run(b.text)
                run.italic, run.bold = b.style == "italic", b.style == "bold"
            elif b.kind == "label":
                doc.add_paragraph().add_run(b.text).bold = True
            elif b.kind == "subheading":
                doc.add_heading(b.text, 2)
            elif b.kind == "equation":
                pic = _eq_picture(b.text, b.style)
                if pic:
                    stream, width, height = pic
                    doc.add_paragraph().add_run().add_picture(stream, width=Inches(width), height=Inches(height))
                else:
                    run = doc.add_paragraph().add_run(b.text)
                    run.font.name, run.font.size = "Courier New", Pt(10)
            elif b.kind == "bullets":
                for item in b.items:
                    doc.add_paragraph(item, style="List Bullet")
            elif b.kind == "table":
                table = doc.add_table(rows=1, cols=len(b.headers), style="Table Grid")
                for cell, text in zip(table.rows[0].cells, b.headers):
                    cell.text = ""
                    cell.paragraphs[0].add_run(str(text)).bold = True
                for row in b.rows:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = str(text)
                doc.add_paragraph()
            elif b.kind == "status":
                for ok, label, detail in b.rows:
                    p = doc.add_paragraph()
                    mark = p.add_run("✔ " if ok else "✖ ")
                    mark.font.color.rgb = RGBColor(0x1A, 0x7F, 0x45) if ok else RGBColor(0xC0, 0x39, 0x2B)
                    p.add_run(f"{label}: ").bold = True
                    p.add_run(str(detail))
            elif b.kind == "kv":
                for k, v in b.rows:
                    p = doc.add_paragraph()
                    p.add_run(f"{k}: ").bold = True
                    p.add_run(str(v))
            elif b.kind == "image" and b.png:
                doc.add_paragraph().add_run(b.label_text).bold = True
                doc.add_picture(io.BytesIO(b.png), width=Inches(6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph().add_run(b.text)
                cap.italic, cap.font.size = True, Pt(9)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- PowerPoint

_SLIDE_W, _SLIDE_H = 13.333, 7.5
_BODY_TOP, _BODY_BOTTOM, _MARGIN = 1.35, 7.0, 0.6
_INK = PptxRGB(0x1D, 0x22, 0x30)
_MUTED = PptxRGB(0x66, 0x6B, 0x7A)


def _text_height(text: str, size_pt: float, width_in: float) -> float:
    chars_per_line = max(10, int(width_in * 72 / (size_pt * 0.52)))
    lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\n"))
    return lines * size_pt * 1.3 / 72 + 0.08


class _Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = PInches(_SLIDE_W), PInches(_SLIDE_H)
        self.slide: Any = None
        self.y = _BODY_TOP
        self.title = ""

    def new_slide(self, title: str, cont: bool = False):
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])      # blank layout
        self.title = title
        shown = f"{title} (cont.)" if cont else title
        self._text(shown, _MARGIN, 0.35, _SLIDE_W - 2 * _MARGIN, 0.8, 30, bold=True)
        self.y = _BODY_TOP
        return self.slide

    def _text(self, text, x, y, w, h, size, bold=False, italic=False, color=_INK, font=None):
        box = self.slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
        tf = box.text_frame
        tf.word_wrap = True
        run = tf.paragraphs[0].add_run()
        run.text = str(text)
        run.font.size, run.font.bold, run.font.italic = PPt(size), bold, italic
        run.font.color.rgb = color
        if font:
            run.font.name = font
        return box

    def room_for(self, h: float) -> bool:
        return self.y + h <= _BODY_BOTTOM

    def ensure(self, h: float) -> None:
        if not self.room_for(h) and self.y > _BODY_TOP:
            self.new_slide(self.title, cont=True)

    def text(self, text, size=16, bold=False, italic=False, color=_INK, font=None) -> None:
        width = _SLIDE_W - 2 * _MARGIN
        h = _text_height(text, size, width)
        self.ensure(h)
        self._text(text, _MARGIN, self.y, width, h, size, bold, italic, color, font)
        self.y += h

    def bullets(self, items: list[str], size=15) -> None:
        for item in items:
            self.text(f"•  {item}", size=size)

    def equation(self, latex: str, style: str, scale: float = 1.35) -> None:
        pic = _eq_picture(latex, style)
        if pic is None:
            self.text(latex, size=13, font="Courier New")
            return
        stream, w, h = pic
        h, w = h * scale, w * scale                     # slides are read from a distance: larger than on paper
        if w > _SLIDE_W - 2 * _MARGIN:
            h, w = h * (_SLIDE_W - 2 * _MARGIN) / w, _SLIDE_W - 2 * _MARGIN
        self.ensure(h + 0.1)
        self.slide.shapes.add_picture(stream, PInches(_MARGIN), PInches(self.y), PInches(w), PInches(h))
        self.y += h + 0.1

    def table(self, headers: list[str], rows: list[list]) -> None:
        size, row_h = 12, 0.36
        remaining = list(rows)
        while True:
            fit = max(1, int((_BODY_BOTTOM - self.y) / row_h) - 1)
            chunk, remaining = remaining[:fit], remaining[fit:]
            n_rows = len(chunk) + 1
            shape = self.slide.shapes.add_table(n_rows, len(headers), PInches(_MARGIN), PInches(self.y),
                                                PInches(_SLIDE_W - 2 * _MARGIN), PInches(row_h * n_rows))
            for c, text in enumerate(headers):
                cell = shape.table.cell(0, c)
                cell.text = str(text)
                cell.text_frame.paragraphs[0].runs[0].font.size = PPt(size)
                cell.text_frame.paragraphs[0].runs[0].font.bold = True
            for r, row in enumerate(chunk, start=1):
                for c, text in enumerate(row):
                    cell = shape.table.cell(r, c)
                    cell.text = str(text)
                    cell.text_frame.paragraphs[0].runs[0].font.size = PPt(size)
            self.y += row_h * n_rows + 0.15
            if not remaining:
                return
            self.new_slide(self.title, cont=True)

    def picture(self, png: bytes, caption: str) -> None:
        try:
            w_px, h_px = Image.open(io.BytesIO(png)).size
        except Exception:  # noqa: BLE001
            return
        max_w, max_h = _SLIDE_W - 2 * _MARGIN, _BODY_BOTTOM - self.y - 0.6
        scale = min(max_w / w_px, max_h / h_px)
        w, h = w_px * scale, h_px * scale
        self.slide.shapes.add_picture(io.BytesIO(png), PInches((_SLIDE_W - w) / 2), PInches(self.y), PInches(w), PInches(h))
        self.y += h + 0.1
        self.text(caption, size=12, italic=True, color=_MUTED)

    def notes(self, text: str) -> None:
        if text:
            self.slide.notes_slide.notes_text_frame.text = text


def _status_lines(rows) -> list[str]:
    return [f"{'✔' if ok else '✖'}  {label}: {detail}" for ok, label, detail in rows]


def _steps_to_slides(deck: _Deck, sec: Section) -> None:
    """One slide per step. Blocks arrive as: subheading (the target), then per step label, equation, [text]."""
    totals: dict[str, int] = {}
    current = ""
    for b in sec.blocks:                                   # first pass: how many steps does each target have
        if b.kind == "subheading":
            current = b.text.removeprefix("Solving for ").strip()
        elif b.kind == "label" and b.text.startswith("Step "):
            totals[current] = totals.get(current, 0) + 1
    target, index, on_step_slide = "", 0, False
    for b in sec.blocks:
        if b.kind == "subheading":
            target, index, on_step_slide = b.text.removeprefix("Solving for ").strip(), 0, False
        elif b.kind == "label" and b.text.startswith("Step "):
            index += 1
            description = b.text.split(":", 1)[1].strip() if ":" in b.text else b.text
            deck.new_slide(f"Solving for {target}: step {index} of {totals.get(target, index)}")
            deck.text(description, size=24, bold=True)
            deck.y += 0.3
            on_step_slide = True
        elif b.kind == "equation" and on_step_slide:
            deck.equation(b.text, "plain", scale=3.2)       # this slide is about the equation: make it big
            deck.y += 0.3
        elif b.kind == "text" and on_step_slide:
            deck.text(b.text, size=18, italic=True, color=_MUTED)
            deck.notes(b.text)


def render_pptx(sections: list[Section], title: str = REPORT_TITLE) -> bytes:
    deck = _Deck()
    deck.new_slide(title)
    deck.y = 3.0
    deck.text(generated_stamp(), size=18, color=_MUTED)
    first = next((s for s in sections if s.id == "problem"), None)
    if first and first.blocks and first.blocks[0].kind == "text":
        deck.text(first.blocks[0].text[:400], size=20)

    for sec in sections:
        if sec.id == "steps":
            _steps_to_slides(deck, sec)
            continue
        deck.new_slide(sec.title)
        for b in sec.blocks:
            if b.kind == "text":
                deck.text(b.text, size=16, bold=b.style == "bold", italic=b.style == "italic")
            elif b.kind == "label":
                deck.text(b.text, size=16, bold=True)
            elif b.kind == "subheading":
                deck.text(b.text, size=18, bold=True)
            elif b.kind == "equation":
                deck.equation(b.text, b.style)
            elif b.kind == "bullets":
                deck.bullets(b.items)
            elif b.kind == "table":
                deck.table(b.headers, b.rows)
            elif b.kind == "status":
                deck.bullets(_status_lines(b.rows), size=13)
            elif b.kind == "kv":
                deck.bullets([f"{k}: {v}" for k, v in b.rows], size=14)
            elif b.kind == "image" and b.png:
                deck.new_slide(b.label_text or "Plot")
                deck.picture(b.png, b.text)
    out = io.BytesIO()
    deck.prs.save(out)
    return out.getvalue()
