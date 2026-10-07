"""Each export format rendered from the shared content: Markdown, PDF, HTML, LaTeX, Word, PowerPoint."""
import io
import re
import shutil
import subprocess
import zipfile
from dataclasses import replace

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from modules import exporter
from modules.report_content import PRESETS, SECTION_TITLES, build_sections
from modules.report_export import FORMAT_KEYS, FORMATS, expected_file, render_export, available_sections
from modules.report_html import count_live_figures, latex_to_svg, render_html
from modules.report_office import render_docx, render_pptx
from modules.report_tex import looks_balanced, package_tex, render_tex, tex_text
from tests.report_helpers import PNG, full_extras, kinematics_ctx, linear_system_ctx, snapshot


def secs(ctx=None, extras=None, include=None, live=False):
    ctx = ctx or kinematics_ctx()
    ex = extras if extras is not None else ctx.extras
    return build_sections(ctx.problem_text, ctx.model, ctx.report, ctx.steps_by_target, ctx.scenarios,
                          ctx.snapshots, replace(ex, live_figures=live), include)


# ------------------------------------------------------------------ the entry point

@pytest.mark.parametrize("fmt", FORMAT_KEYS)
def test_every_format_renders_and_has_the_right_name_and_mime(fmt):
    ctx = kinematics_ctx(extras=full_extras())
    out = render_export(fmt, ctx, PRESETS["Student handout"])
    spec = next(f for f in FORMATS if f.key == fmt)
    assert out.name == "solved_problem" + spec.ext and out.mime == spec.mime
    assert len(out.data) > 500


def test_an_unknown_format_or_section_is_refused():
    ctx = kinematics_ctx()
    with pytest.raises(ValueError, match="format"):
        render_export("rtf", ctx)
    with pytest.raises(ValueError, match="section"):
        render_export("md", ctx, ["problem", "bogus"])
    with pytest.raises(ValueError, match="format"):
        expected_file("rtf", ctx)


def test_choosing_fewer_sections_gives_a_smaller_report_in_every_format():
    ctx = kinematics_ctx(extras=full_extras())
    for fmt in ("md", "html", "tex", "docx", "pptx", "pdf"):
        full = render_export(fmt, ctx, PRESETS["Full audit"]).data
        small = render_export(fmt, ctx, PRESETS["Quick summary"]).data
        assert len(small) < len(full), fmt


def test_available_sections_lists_only_those_with_content():
    ctx = kinematics_ctx()
    got = available_sections(ctx)
    assert "matrix" not in got and "plots" not in got and "steps" in got
    assert "matrix" in available_sections(linear_system_ctx())
    assert "plots" in available_sections(kinematics_ctx(snapshots=[snapshot()]))


def test_the_expected_name_is_known_before_building_and_matches_what_is_built():
    plain, with_plot = kinematics_ctx(), kinematics_ctx(snapshots=[snapshot()])
    assert expected_file("tex", plain)[0] == "solved_problem.tex"
    assert expected_file("tex", with_plot)[0] == "solved_problem_tex.zip"
    assert expected_file("tex", with_plot, ["problem"])[0] == "solved_problem.tex"       # plots not chosen
    for ctx, include in ((plain, None), (with_plot, None), (with_plot, ["problem"])):
        name, mime = expected_file("tex", ctx, include)
        built = render_export("tex", ctx, include)
        assert (built.name, built.mime) == (name, mime)


# ------------------------------------------------------------------ Markdown and PDF (as before)

def test_markdown_keeps_the_headings_it_always_had_and_adds_the_results_table():
    md = exporter.build_markdown("p", *_args(kinematics_ctx()))
    for h in ("## Problem", "## Derived equations", "## Variables", "## Verification detail", "## Step-by-step solution",
              "## Confidence report", "## Results"):
        assert h in md
    assert "### Solving for `a`" in md and "$$ v = a t + u $$" in md and "| a | 2 | m/s^2 |" in md


def _args(ctx):
    return ctx.model, ctx.report, ctx.steps_by_target, ctx.scenarios


def test_markdown_table_cells_cannot_break_the_table():
    ctx = kinematics_ctx()
    ctx.model.variables[0].meaning = "has | a pipe\nand a newline"
    md = exporter.build_markdown("p", *_args(ctx))
    row = next(l for l in md.splitlines() if l.startswith("| v |"))
    assert "has \\| a pipe and a newline" in row and row.count("|") == 5 + 1


def test_the_build_functions_still_take_include_and_extras():
    ctx = kinematics_ctx(extras=full_extras())
    md = exporter.build_markdown("p", *_args(ctx), extras=ctx.extras, include=["followups", "provenance"])
    assert "## Follow-up questions" in md and "## Reproducibility" in md and "## Variables" not in md
    pdf = exporter.build_pdf_bytes("p", *_args(ctx), extras=ctx.extras, include=["problem", "results"])
    assert pdf[:5] == b"%PDF-"


def test_pdf_text_contains_the_chosen_sections_only():
    ctx = kinematics_ctx(extras=full_extras())
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(render_export("pdf", ctx, ["problem", "followups"]).data)).pages)
    assert "Follow-up questions" in text and "what if t doubles?" in text
    assert "Variables" not in text and "Step-by-step" not in text


# ------------------------------------------------------------------ HTML

def _outside_scripts(html: str) -> str:
    return re.sub(r"<script>.*?</script>", "", html, flags=re.S)


def test_the_html_report_is_self_contained():
    ctx = kinematics_ctx(extras=full_extras(), snapshots=[snapshot()])
    page = render_export("html", ctx).data.decode()
    body = _outside_scripts(page)
    assert not re.search(r'(src|href)=["\']https?:', body), "no external resources"
    assert "@import" not in page and "<link" not in page
    assert not re.search(r"url\(\s*['\"]?https?:", page)
    assert page.startswith("<!DOCTYPE html>") and "</html>" in page


def test_equations_are_inline_svg_that_follow_the_text_colour():
    page = render_html(secs())
    assert page.count('class="eq"') >= 4
    assert "<svg fill=\"currentColor\"" in page and "<metadata>" not in page
    assert 'data-latex="v = a t + u"' in page                        # click-to-copy carries the LaTeX


def test_an_equation_mathtext_cannot_typeset_falls_back_to_its_source_not_a_crash():
    assert latex_to_svg(r"\unknowncommand{x}") is None
    ctx = kinematics_ctx()
    s = secs(ctx)
    next(b for sec in s if sec.id == "equations" for b in sec.blocks if b.kind == "equation").text = r"\unknowncommand{x}"
    assert "<code>\\unknowncommand{x}</code>" in render_html(s)


def test_a_raw_unparsed_equation_is_shown_as_code():
    ctx = kinematics_ctx()
    ctx.model.equations[0].sympy_eq = None
    page = render_html(secs(ctx))
    assert "<code>" in page and ctx.model.equations[0].raw_expression.replace("&", "&amp;") in page.replace("&#x27;", "'")


def test_text_from_the_user_and_the_model_is_escaped():
    evil = '<script>alert(1)</script><img src=x onerror=alert(2)>"'
    ctx = kinematics_ctx(problem_text=evil, extras=full_extras())
    ctx.extras.followups[0].answer = evil
    ctx.model.variables[0].meaning = evil
    ctx.steps_by_target["a"][0].explanation = evil
    page = render_export("html", ctx).data.decode()
    body = _outside_scripts(page)
    assert "<script>alert(1)" not in body and "<img src=x" not in body
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in body


def test_without_live_plots_there_is_no_javascript_library_at_all():
    ctx = kinematics_ctx()
    ctx.extras.include_provenance = False
    page = render_html(secs(ctx))
    assert "Plotly" not in page and len(page) < 200_000
    assert count_live_figures(page) == 0


def test_with_live_plots_the_figures_are_drawn_in_the_page():
    ctx = kinematics_ctx(extras=full_extras())
    page = render_export("html", ctx).data.decode()
    assert count_live_figures(page) == 3 and "Plotly.newPlot" in page
    assert len(page) > 1_000_000                                       # the library itself is embedded


def test_a_marked_plot_is_live_when_it_has_a_figure_and_an_image_when_it_does_not():
    live = snapshot("Live one", "c", figure_json='{"data":[{"type":"scatter","x":[1,2],"y":[1,4]}],"layout":{}}')
    static = snapshot("Static one", "c")
    ctx = kinematics_ctx(snapshots=[live, static])
    page = render_export("html", ctx, ["plots"]).data.decode()
    assert count_live_figures(page) == 1
    assert _outside_scripts(page).count("data:image/png;base64") == 1 and 'alt="Static one"' in page


def test_figure_json_cannot_close_the_script_tag():
    evil = '{"data":[],"layout":{"title":{"text":"</script><b>x"}}}'
    page = render_html(secs(kinematics_ctx(snapshots=[snapshot("e", "c", figure_json=evil)])))
    assert "</script><b>x" not in page and "<\\/script><b>x" in page


def test_the_navigation_links_every_section():
    page = render_html(secs(extras=full_extras()))
    for s in secs(extras=full_extras()):
        assert f'href="#sec-{s.id}"' in page and f'id="sec-{s.id}"' in page


def test_the_status_marks_and_dark_mode_and_print_styles_are_present():
    page = render_html(secs())
    assert "✔" in page and "prefers-color-scheme:dark" in page and "@media print" in page


# ------------------------------------------------------------------ LaTeX

def test_the_latex_source_is_balanced_and_complete():
    src, files = render_tex(secs(extras=full_extras()))
    assert looks_balanced(src) and src.startswith("\\documentclass") and src.rstrip().endswith("\\end{document}")
    assert files == {}


@pytest.mark.parametrize("text,expected", [("50% & more_x", r"50\% \& more\_x"), ("a → b", r"a $\to$ b"),
                                           ("±5", r"$\pm$5"), ("✅ ok", "[PASS] ok"), ("a—b", "a---b")])
def test_latex_prose_is_escaped_and_unicode_is_mapped(text, expected):
    assert tex_text(text) == expected


def test_a_raw_equation_is_set_as_text_not_math():
    ctx = kinematics_ctx()
    ctx.model.equations[0].sympy_eq = None
    src, _ = render_tex(secs(ctx))
    assert "\\texttt{" in src and looks_balanced(src)


def test_plots_make_the_latex_a_zip_with_the_figures_it_refers_to():
    ctx = kinematics_ctx(snapshots=[snapshot("P", "cap")])
    sections = secs(ctx)
    src, files = render_tex(sections)
    assert files == {"figures/plot-1.png": PNG} and "\\includegraphics" in src and "figures/plot-1.png" in src
    name, data, mime = package_tex(sections)
    assert (name, mime) == ("report_tex.zip", "application/zip")
    assert sorted(zipfile.ZipFile(io.BytesIO(data)).namelist()) == ["figures/plot-1.png", "report.tex"]


@pytest.mark.skipif(shutil.which("pdflatex") is None, reason="no TeX installation")
def test_the_latex_actually_compiles(tmp_path):
    ctx = kinematics_ctx(extras=full_extras(), snapshots=[snapshot("P", "cap")])
    name, data, _ = package_tex(secs(ctx))
    zipfile.ZipFile(io.BytesIO(data)).extractall(tmp_path)
    run = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "report.tex"], cwd=tmp_path,
                         capture_output=True, text=True, timeout=120)
    assert run.returncode == 0, run.stdout[-1500:]
    assert (tmp_path / "report.pdf").stat().st_size > 5000


# ------------------------------------------------------------------ Word

def test_the_word_document_has_a_heading_per_section_real_tables_and_equation_pictures():
    sections = secs(extras=full_extras())
    doc = Document(io.BytesIO(render_docx(sections)))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == [s.title for s in sections]
    n_tables = sum(1 for s in sections for b in s.blocks if b.kind == "table")
    assert len(doc.tables) == n_tables
    n_equations = sum(1 for s in sections for b in s.blocks if b.kind == "equation")
    assert len(doc.inline_shapes) == n_equations


def test_a_marked_plot_is_a_picture_with_its_caption_in_word():
    ctx = kinematics_ctx(snapshots=[snapshot("My plot", "its caption")])
    doc = Document(io.BytesIO(render_export("docx", ctx, ["plots"]).data))
    assert any(p.text == "My plot" for p in doc.paragraphs) and any(p.text == "its caption" for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1


def test_word_status_lines_carry_the_pass_or_fail_mark():
    doc = Document(io.BytesIO(render_export("docx", kinematics_ctx(), ["verification"]).data))
    assert any(p.text.startswith("✔ ") for p in doc.paragraphs)


# ------------------------------------------------------------------ PowerPoint

def test_the_deck_has_one_slide_per_step_titled_with_its_target_and_position():
    ctx = kinematics_ctx()
    prs = Presentation(io.BytesIO(render_export("pptx", ctx, ["steps"]).data))
    titles = [next(iter(s.shapes)).text_frame.text for s in prs.slides]
    assert titles[0] == "Math Representation System -- Solved Problem"
    step_titles = titles[1:]
    assert len(step_titles) == 14
    assert step_titles[0] == "Solving for a: step 1 of 7" and step_titles[7] == "Solving for d: step 1 of 7"
    assert step_titles[-1] == "Solving for d: step 7 of 7"


def test_each_step_slide_has_the_equation_as_a_picture_and_the_explanation_in_the_notes():
    ctx = kinematics_ctx()
    for steps in ctx.steps_by_target.values():
        steps[0].explanation = "Why this step: it sets up the equation."
    prs = Presentation(io.BytesIO(render_export("pptx", ctx, ["steps"]).data))
    step = prs.slides[1]
    assert any(sh.shape_type == 13 for sh in step.shapes)                    # a picture
    with_notes = [s for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text]
    explained = sum(1 for steps in ctx.steps_by_target.values() for st in steps if st.explanation)
    assert len(with_notes) == explained and explained > 0


def test_other_sections_get_a_slide_each_and_tables_are_real_tables():
    prs = Presentation(io.BytesIO(render_export("pptx", kinematics_ctx(), ["problem", "results", "variables"]).data))
    titles = [next(iter(s.shapes)).text_frame.text for s in prs.slides]
    assert titles[1:] == ["Problem", "Results", "Variables"]
    assert any(sh.has_table for sh in prs.slides[2].shapes)


def test_a_long_table_continues_on_the_next_slide_instead_of_overflowing():
    ctx = kinematics_ctx()
    ctx.model.variables = ctx.model.variables * 6                               # 30 rows
    prs = Presentation(io.BytesIO(render_export("pptx", ctx, ["variables"]).data))
    titles = [next(iter(s.shapes)).text_frame.text for s in prs.slides]
    assert any("(cont.)" in t for t in titles)
    for slide in prs.slides:                                                    # nothing hangs below the slide
        for sh in slide.shapes:
            assert sh.top + sh.height <= prs.slide_height * 1.02


def test_a_marked_plot_gets_a_full_slide_with_its_caption():
    ctx = kinematics_ctx(snapshots=[snapshot("Sweep result", "a caption")])
    prs = Presentation(io.BytesIO(render_export("pptx", ctx, ["plots"]).data))
    last = prs.slides[len(prs.slides) - 1]
    assert next(iter(last.shapes)).text_frame.text == "Sweep result"
    assert any(sh.shape_type == 13 for sh in last.shapes)
    assert any(sh.has_text_frame and sh.text_frame.text == "a caption" for sh in last.shapes)


def test_every_section_title_is_findable_in_the_deck_text():
    sections = secs(extras=full_extras(), include=[s for s in SECTION_TITLES if s != "steps"])
    prs = Presentation(io.BytesIO(render_pptx(sections)))
    deck_titles = {next(iter(s.shapes)).text_frame.text.replace(" (cont.)", "") for s in prs.slides}
    assert {s.title for s in sections} <= deck_titles
